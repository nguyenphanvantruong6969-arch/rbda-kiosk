"""Xuất kết quả: tệp tổng, tệp tổng hợp, tệp theo CLB/buổi và tệp Excel.

Một phần của PipelineAPI (xem api.py) — tách riêng cho dễ đọc; mọi hàm
public ở đây vẫn là hàm của PipelineAPI và giao diện gọi được như cũ.
"""

import datetime
import os
import sys

from api_chung import (
    _an_toan_cho_excel,
    _ghi_csv,
    _now,
    _phan_tram_vn,
    _so_vn,
    duong_dan_khong_de,
    thu_muc_tai_ve,
)
from i18n_errors import err
from i18n_errors import phan_hoi_loi as _fail
from i18n_errors import phan_hoi_ok as _ok
from rbda_priority_pipeline import BUOI_MAC_DINH, dau_van_tay_du_lieu


class XuatKetQuaMixin:
    # -----------------------------------------------------------------
    # XUẤT KẾT QUẢ
    #
    # File này là SẢN PHẨM CUỐI của cả quy trình — nhà trường dán bảng,
    # gửi phụ huynh, phát cho giáo viên phụ trách từng CLB. Bản cũ chỉ
    # xuất hai cột mã (student_id,club_id): nhìn vào không biết em nào
    # tên gì, đỗ CLB nào, đỗ nguyện vọng thứ mấy. Toàn bộ dữ liệu đó ĐÃ
    # có sẵn trong DB, chỉ là câu lệnh xuất không lấy.
    # -----------------------------------------------------------------

    # Ký tự KHÔNG được phép có trong tên file trên Windows, cộng thêm
    # dấu phân cách đường dẫn. club_id do trường tự đặt và
    # create_or_update_club KHÔNG giới hạn ký tự, nên một mã như
    # "../ngoai" hay "khoi 10/11" sẽ làm file rơi ra ngoài thư mục kết
    # quả (hoặc ghi đè file khác) nếu ghép thẳng vào tên file.
    _KY_TU_CAM_TRONG_TEN_FILE = '<>:"/\\|?*'

    @staticmethod
    def _ten_file_an_toan(raw: str, mac_dinh: str = "club") -> str:
        ten = "".join(
            "_" if (c in XuatKetQuaMixin._KY_TU_CAM_TRONG_TEN_FILE or ord(c) < 32) else c
            for c in (raw or "")
        ).strip(" .")
        # "..", "." và chuỗi rỗng đều không dùng làm tên file được
        if not ten or ten in (".", ".."):
            return mac_dinh
        # Windows không cho tạo tệp tên CON, NUL, COM1... (kể cả khi có
        # đuôi: "nul.csv" cũng hỏng). CLB mã "CON" từng làm lần xuất báo lỗi.
        if ten.split(".")[0].upper() in XuatKetQuaMixin._TEN_DANH_RIENG_WINDOWS:
            return "_" + ten
        return ten

    _TEN_DANH_RIENG_WINDOWS = frozenset(
        ["CON", "PRN", "AUX", "NUL"]
        + ["COM%d" % i for i in range(1, 10)]
        + ["LPT%d" % i for i in range(1, 10)])

    @staticmethod
    def _ten_chua_dung(ten: str, da_dung: set) -> str:
        """`ten`, hoặc `ten_2`, `ten_3`... — cái đầu tiên CHƯA dùng.

        So KHÔNG phân biệt hoa/thường: ổ NTFS trên Windows coi `clb_a.csv`
        và `CLB_A.csv` là một tệp, mà phần mềm lại cho phép hai mã CLB đó
        cùng tồn tại (_khop_club_id) — tệp sau từng ghi đè tệp trước, một
        giáo viên cầm nhầm danh sách. Thử đến khi trống thật, vì cách cũ
        (ghép số đếm) vẫn đụng một CLB tên sẵn là `a_2`.
        """
        ung_vien, n = ten, 1
        while ung_vien.casefold() in da_dung:
            n += 1
            ung_vien = f"{ten}_{n}"
        da_dung.add(ung_vien.casefold())
        return ung_vien

    @staticmethod
    def _xuat_excel(path: str, cac_bang: list) -> None:
        """Ghi mọi bảng vào MỘT tệp Excel, mỗi bảng một trang.

        VÌ SAO CẦN, KHI ĐÃ CÓ CSV. Một lần xuất đẻ ra năm tệp `.csv` và hai
        thư mục. Người nhận phải mở từng tệp, và Excel mở `.csv` thì mất hết
        định dạng — cột không giãn, dòng tiêu đề trôi khi cuộn. Tệp này là
        thứ mang đi họp; mấy tệp `.csv` vẫn còn nguyên cho ai cần đổ vào
        phần mềm khác.

        Chỉ gom các bảng của CẢ TRƯỜNG. Bảng từng câu lạc bộ vẫn nằm ở thư
        mục `.csv` riêng, vì một trường 20 câu lạc bộ sẽ thành 20 trang —
        không ai dùng nổi một workbook như thế, mà từng tệp rời lại đúng cái
        giáo viên phụ trách cần cầm đi.

        `cac_bang`: [(tên trang, [tiêu đề cột], [[ô, ...], ...]), ...].
        Bảng nào rỗng thì BỎ QUA — một trang trống chỉ làm người đọc tưởng
        dữ liệu bị mất.
        """
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        wb.remove(wb.active)

        DAM = Font(bold=True)
        NEN = PatternFill("solid", fgColor="DDE5F0")

        for ten, cot, dong in cac_bang:
            if not dong and not cot:
                continue
            # Ten trang Excel: toi da 31 ky tu, va cam : \ / ? * [ ]
            ten_sach = "".join(
                "-" if c in ':\\/?*[]' else c for c in (ten or "Trang"))[:31]
            ws = wb.create_sheet(ten_sach)
            if cot:
                ws.append(list(cot))
                for o in ws[1]:
                    o.font = DAM
                    o.fill = NEN
                    o.alignment = Alignment(vertical="center", wrap_text=True)
                # Khoa dong tieu de: cuon xuong dong 300 van con biet cot nao
                # la cot nao. Day dung la ly do nguoi ta ghet doc .csv dai.
                ws.freeze_panes = "A2"
            for d in dong:
                ws.append(list(d))
                # Chan CONG THUC, giong ban CSV (an_toan_cho_excel trong
                # export_csv). openpyxl coi moi chuoi bat dau bang "=" la
                # cong thuc: hoc sinh ten "=HYPERLINK(...)" se thanh mot
                # lien ket song trong tep giao vien mo. Ep ve kieu chuoi
                # thi Excel hien NGUYEN VAN ten, khong tinh, khong them
                # dau nhay nhu ban CSV phai lam.
                for o in ws[ws.max_row]:
                    if o.data_type == "f":
                        o.data_type = "s"
                        o.quotePrefix = True

            # Gian cot theo noi dung that, chan tren 60 de mot ten dai khong
            # keo cot rong ra ca man hinh.
            for i in range(1, (len(cot) or 1) + 1):
                rong = len(str(cot[i - 1])) if cot else 8
                for d in dong[:400]:      # 400 dong dau la du de uoc luong
                    if i - 1 < len(d):
                        rong = max(rong, len(str(d[i - 1] if d[i - 1] is not None else "")))
                ws.column_dimensions[get_column_letter(i)].width = min(rong + 2, 60)

        wb.save(path)

    def _bang_bao_cao(self, so: bool) -> dict:
        """Bốn bảng báo cáo cả trường, dựng MỘT lần cho cả hai nơi dùng.

        Tệp tổng hợp (.csv, đọc bằng mắt) và các trang Excel (để tính tiếp)
        trước đây mỗi nơi tự dựng lại cùng bốn bảng — cùng cột, cùng phép
        tính, chỉ khác cách viết số. Sửa một nơi mà quên nơi kia là hai tệp
        của CÙNG một lần xuất nói hai điều khác nhau.

        `so=False`: số viết kiểu Việt để đọc ("84,5%", "(không có chỗ)").
        `so=True` : số thô để Excel tính được (84.5, ô trống).

        Trả về {tên: (số tóm tắt cho tiêu đề, [cột], [[ô]])}, chỉ gồm bảng
        đọc được dữ liệu.
        """
        def ti_le_choi(x):
            if x is None:
                return "" if so else "(không có chỗ)"
            return x if so else _so_vn(x)

        def phan_tram(a, b):
            if so:
                return round(100 * a / b, 1) if b else ""
            return _phan_tram_vn(a, b)

        def khai(e):
            return "; ".join(
                "%s (%s)" % (k["ten_clb"] or k["club_id"],
                             (k["ti_le_choi"] if so else _so_vn(k["ti_le_choi"]))
                             if k["ti_le_choi"] is not None else "không có chỗ")
                for k in e["da_khai"]) or "(chưa khai nguyện vọng nào)"

        ra = {}
        lap = self.get_club_fill_stats()
        if lap.get("ok"):
            ra["clb"] = (len(lap["data"]), [
                "Buổi", "Mã CLB", "Tên CLB", "Số em đăng ký",
                "Đặt nguyện vọng 1", "Tỉ lệ chọi", "Đã xếp", "Sức chứa",
                "Tỉ lệ lấp đầy" + (" %" if so else ""),
                "Trong đó diện dự trữ", "Chỉ tiêu dự trữ"],
                [[c.get("buoi"), c.get("club_id"), c.get("name"),
                  c.get("so_dang_ky"), c.get("so_dat_nv1"),
                  ti_le_choi(c.get("ti_le_choi")),
                  c.get("matched"), c.get("capacity") or 0,
                  phan_tram(c.get("matched") or 0, c.get("capacity") or 0),
                  c.get("matched_reserve"), c.get("reserve_capacity")]
                 for c in lap["data"]])

        nv = self.get_phan_bo_nguyen_vong()
        if nv.get("ok") and nv["data"]["tong_da_xep"]:
            n = nv["data"]
            tong = n["tong_da_xep"]
            dong = [[("Nguyện vọng %d trở lên" % h["hang"]) if h["gop"]
                     else ("Nguyện vọng %d" % h["hang"]),
                     h["so_em"], phan_tram(h["so_em"], tong)]
                    for h in n["phan_bo"]]
            if n["khong_ro"]:
                dong.append(["Không rõ thứ hạng", n["khong_ro"],
                             phan_tram(n["khong_ro"], tong)])
            ra["nv"] = (tong, ["Nguyện vọng", "Số chỗ",
                               "Tỉ lệ %" if so else "Tỉ lệ"], dong)

        dt = self.get_suat_du_tru()
        if dt.get("ok") and dt["data"]:
            ra["dt"] = (sum(c["con_thua"] for c in dt["data"]), [
                "Buổi", "Mã CLB", "Tên CLB", "Nhãn dự trữ",
                "Chỉ tiêu dự trữ", "Em thuộc nhóm có đăng ký",
                "Đã dùng", "Còn thừa"],
                [[c["buoi"], c["club_id"], c["name"], c["reserve_group"],
                  c["reserve_capacity"], c["nhom_dang_ky"], c["da_dung"],
                  c["con_thua"]] for c in dt["data"]])

        tt = self.get_em_chua_co_cho()
        if tt.get("ok"):
            d = tt["data"]
            ra["tt"] = (d["so_em"], [
                "Mã học sinh", "Họ tên", "Số nguyện vọng đã khai",
                "Đã khai những câu lạc bộ (kèm tỉ lệ chọi)"],
                [[e["student_id"], e["name"] or "", e["so_nguyen_vong"], khai(e)]
                 for e in d["danh_sach"]])
        return ra

    def _bang_tong_hop(self) -> list[list]:
        """Các dòng của tệp tổng hợp — dùng chung cho bản CSV và bản Excel.

        VÌ SAO TỆP NÀY PHẢI CÓ. Mọi tệp xuất ra trước đây đều nói *ai vào câu
        lạc bộ nào*, không tệp nào nói *kết quả này đến từ đâu*. Hai lần chạy
        khác hạt giống cho ra hai bộ tệp trông y hệt nhau: cùng tên cột, cùng
        số dòng, cùng hình dạng. Trường in ra dán bảng rồi sáu tháng sau bị
        hỏi "bảng này tính thế nào" thì không còn gì để trả lời.

        Mà cả thiết kế của phần mềm dựa trên một lời hứa: *ai cũng tính lại
        được*. Lời hứa ấy cần đúng ba thứ — bộ số bốc thăm đã khoá, hạt
        giống, và cách bốc thăm — và trước bản vá này **không thứ nào** đi
        theo tệp kết quả ra khỏi phần mềm.

        Bảng còn gộp mấy con số vốn chỉ xem được trên màn hình: độ phủ, tỉ lệ
        lấp đầy từng câu lạc bộ, tải từng buổi, và danh sách em chưa có chỗ
        nào. Đó đúng là những thứ nhà trường phải xử lý tiếp, nên chúng cần
        nằm trong tệp mang đi họp chứ không nằm trong một tab phải mở phần
        mềm mới thấy.
        """
        def phan(tieu_de):
            return [[], [tieu_de]]

        ra: list[list] = [["TỔNG HỢP KẾT QUẢ PHÂN BỔ CÂU LẠC BỘ (RB-DA)"]]

        # ---- 1. Dau vet lan chay ----
        meta = self.get_last_run_info()
        m = meta.get("data") if meta.get("ok") else None
        khoa = self.get_stb_lock_status()
        k = khoa.get("data") if khoa.get("ok") else {}

        ra += phan("DẤU VẾT LẦN CHẠY — ba dòng đầu là thứ cần để tính lại")
        ra.append(["Mục", "Giá trị"])
        if m:
            ra += [
                ["Hạt giống bốc thăm", m.get("seed")],
                ["Cách bốc thăm", m.get("che_do_boc_tham")],
                ["Bộ số bốc thăm khoá lúc", (k or {}).get("locked_at") or "(chưa khoá)"],
                ["Thời điểm chạy", m.get("run_at")],
                ["Số vòng xử lý", m.get("rounds_run")],
                ["Số buổi lần chạy này phủ", m.get("so_buoi")],
                ["Buổi đã chạy", (m.get("buoi_da_chay") or "").replace(",", ", ")],
            ]
        else:
            ra.append(["(chưa chạy lần nào)", ""])
        ra.append(["Thời điểm xuất tệp này", _now()])
        # Tep mang di hop phai tu noi duoc no con dung khong. Man hinh co dai
        # canh bao, nhung tep ra khoi phan mem thi chi con dong nay.
        tt = self.get_trang_thai_ket_qua()
        if tt.get("ok") and tt["data"]["co_ket_qua"]:
            ra.append(["Kết quả còn khớp dữ liệu",
                       "CÓ" if not tt["data"]["ket_qua_cu"] else
                       "KHÔNG — dữ liệu đã sửa sau lần chạy này, kết quả có "
                       "thể đã cũ. Hãy chạy lại rồi xuất lại."])

        # ---- 2. Do phu ----
        dp = self.get_do_phu()
        if dp.get("ok"):
            d = dp["data"]
            tong = d["tong_hoc_sinh"] or 0
            ra += phan("ĐỘ PHỦ — bao nhiêu em được mấy câu lạc bộ")
            ra.append(["Số câu lạc bộ trong tuần", "Số em", "Tỉ lệ"])
            for h in d["phan_bo"]:
                ra.append([h["so_clb"], h["so_em"],
                           _phan_tram_vn(h["so_em"], tong)])
            ra.append(["Tổng số học sinh", tong, ""])
            ra.append(["Trung bình số câu lạc bộ mỗi em",
                       _so_vn(d["trung_binh_clb"]), ""])

        bang = self._bang_bao_cao(so=False)

        # ---- 3. Ti le lap day tung CLB ----
        # Cot "So em dang ky" KHONG cong don duoc: trong cung mot buoi, mot em
        # khai ba cau lac bo thi duoc dem o ca ba dong, ma em ay chi lay duoc
        # MOT cho. Cong cot nay roi chia cho tong chi tieu se ra mot ti le
        # choi phong len vo nghia — nen cau canh bao di ngay trong tieu de.
        if "clb" in bang:
            _, cot, dong = bang["clb"]
            ra += phan("TỪNG CÂU LẠC BỘ — \"số em đăng ký\" KHÔNG cộng dồn "
                       "được (một em khai ba câu lạc bộ được đếm ở cả ba dòng)")
            ra.append(cot)
            ra += dong

        # ---- 4. Tai tung buoi ----
        tai = self.get_tai_theo_buoi()
        if tai.get("ok") and len(tai["data"]) > 1:
            ra += phan("TẢI TỪNG BUỔI — đọc để biết buổi nào chật, buổi nào rộng")
            ra.append(["Buổi", "Số CLB", "Tổng chỗ", "Chỉ tiêu dự trữ",
                       "Số học sinh khai nguyện vọng", "Tỉ lệ chọi"])
            for t in tai["data"]:
                ra.append([
                    t.get("buoi"), t.get("so_clb"), t.get("tong_cho"),
                    t.get("tong_du_tru"), t.get("so_hoc_sinh"),
                    # None = khong co cho nao, khac han "thua cho" — xem
                    # chu thich trong get_tai_theo_buoi.
                    _so_vn(t.get("ti_le_choi")) if t.get("ti_le_choi") is not None
                    else "(không có chỗ)",
                ])

        # ---- 5. Nguyen vong thu may thi duoc ----
        # Do phu noi bao nhieu em co cho, khong noi cho ay co phai thu em
        # muon khong. Hai lan chay rat khac nhau van cho cung mot con so do
        # phu — day la con so phan biet duoc chung.
        if "nv" in bang:
            tong, cot, dong = bang["nv"]
            ra += phan("NGUYỆN VỌNG THỨ MẤY THÌ ĐƯỢC — trên %d chỗ đã xếp" % tong)
            ra.append(cot)
            ra += dong

        # ---- 6. Suat du tru dung toi dau ----
        # Mot suat khong ai dung KHONG bao loi va KHONG hien o dau — no lang
        # le thanh mot suat thuong, co che van chay dung, con y dinh uu tien
        # thi boc hoi. Cot cuoi la con so de quyet dinh noi nhan chinh sach
        # hay ha chi tieu du tru nam sau.
        if "dt" in bang:
            thua, cot, dong = bang["dt"]
            ra += phan("SUẤT DỰ TRỮ DÙNG TỚI ĐÂU — còn thừa %d suất" % thua)
            ra.append(cot)
            ra += dong

        # ---- 7. Em chua co CLB nao ----
        # Nhom nha truong phai xu ly tiep, nen no nam trong tep mang di hop.
        # Doc tu `get_em_chua_co_cho` chu khong tu `get_do_phu`: ban kia cat
        # danh sach con 200 em cho giao dien, va chi co TEN — ma cach xu ly
        # phu thuoc vao em da khai may nguyen vong, vao CLB chat den dau.
        if "tt" in bang:
            so_em, cot, dong = bang["tt"]
            ra += phan("HỌC SINH CHƯA CÓ CÂU LẠC BỘ NÀO — %d em" % so_em)
            if dong:
                ra.append(cot)
                ra += dong
            else:
                ra.append(["(không có em nào — mọi em đều có ít nhất một chỗ)", ""])

        return ra
    def _duong_xuat(self, output_path: str, ten_mac_dinh: str) -> str:
        """Chỗ đặt một tệp (hoặc thư mục) xuất ra — MỘT luật cho mọi lần xuất.

        Rỗng hoặc TƯƠNG ĐỐI -> thư mục Tải xuống (lùi về cạnh app.db nếu
        không ghi được), và không ghi đè tệp đã có: `x.csv` thành
        `x (2).csv` như trình duyệt. Đường dẫn TUYỆT ĐỐI là ý muốn rõ ràng
        của bên gọi -> tôn trọng nguyên văn. Xem `export_csv`.
        """
        output_path = (output_path or "").strip() or ten_mac_dinh
        if os.path.isabs(output_path):
            return output_path
        thu_muc = self.thu_muc_xuat or thu_muc_tai_ve()
        if not (thu_muc and os.path.isdir(thu_muc)
                and os.access(thu_muc, os.W_OK)):
            thu_muc = os.path.dirname(os.path.abspath(self.db_path))
        # Chi tranh ghi de o duong CHUONG TRINH tu chon.
        return duong_dan_khong_de(
            os.path.join(thu_muc, os.path.basename(output_path)))
    def _goc_xuat_hop_le(self) -> list:
        """Những thư mục GỐC mà tệp xuất có thể nằm trong — xem `_duong_xuat`."""
        goc = [self.thu_muc_xuat or thu_muc_tai_ve(),
               os.path.dirname(os.path.abspath(self.db_path))]
        return [os.path.realpath(g) for g in goc if g]
    @staticmethod
    def _nam_trong(thu_muc: str, goc: str) -> bool:
        """`thu_muc` là `goc` hoặc nằm bên trong nó.

        `commonpath` ném ValueError khi hai đường dẫn KHÁC Ổ ĐĨA (Windows:
        C:\\ và D:\\). Khác ổ thì chắc chắn không nằm trong nhau — trả False
        để phép so đi tiếp tới gốc hợp lệ kế tiếp, thay vì hỏng cả lần mở
        (tải xuống ở C: không ghi được, tệp rơi về cạnh app.db ở D:).
        """
        try:
            return os.path.commonpath([thu_muc, goc]) == goc
        except ValueError:
            return False
    def mo_thu_muc(self, path: str):
        """Mở thư mục chứa tệp vừa xuất trong trình quản lý tệp của máy.

        Chỉ hiện đường dẫn thì giáo viên phải tự lần tới thư mục Tải xuống
        — việc không phải ai cũng làm được. `path` là một TỆP (mở thư mục
        chứa nó) hoặc một THƯ MỤC (mở chính nó, vd thư mục dữ liệu đầu vào).

        AN TOÀN: hàm này gọi được từ JavaScript, nên CHỈ mở thư mục nằm
        trong chỗ phần mềm tự đặt tệp xuất (`_goc_xuat_hop_le`). So bằng
        `realpath` + `commonpath` để "..", liên kết tượng trưng hay đường
        dẫn lạ đều không lọt ra ngoài. Không bao giờ chạy qua shell.
        """
        try:
            p = os.path.realpath((path or "").strip() or ".")
            thu_muc = p if os.path.isdir(p) else os.path.dirname(p)
            if not os.path.isdir(thu_muc):
                return _fail(err("thu_muc_khong_hop_le", path=path))
            if not any(self._nam_trong(thu_muc, g) for g in self._goc_xuat_hop_le()):
                return _fail(err("thu_muc_khong_hop_le", path=path))
            if sys.platform == "win32":
                os.startfile(thu_muc)  # noqa: S606 — thu muc da duoc soat o tren
            else:
                import subprocess
                lenh = "open" if sys.platform == "darwin" else "xdg-open"
                subprocess.Popen([lenh, thu_muc], stdout=subprocess.DEVNULL,
                                 stderr=subprocess.DEVNULL)
            return _ok({"dir": thu_muc})
        except Exception as e:
            return _fail(err("thu_muc_khong_mo_duoc", detail=str(e)))

    # -----------------------------------------------------------------
    # KẾT QUẢ CÓ CÒN KHỚP DỮ LIỆU KHÔNG — và đã đổi gì so với lần trước
    # -----------------------------------------------------------------
    @staticmethod
    def _dau_van_tay_du_lieu(cur) -> str:
        """Xem `rbda_priority_pipeline.dau_van_tay_du_lieu` — MỘT định nghĩa
        "dữ liệu quyết định kết quả" cho cả lần chạy lẫn phép so sau đó."""
        return dau_van_tay_du_lieu(cur)
    def get_trang_thai_ket_qua(self):
        """Kết quả đang lưu có còn khớp dữ liệu hiện tại không.

        `ket_qua_cu` = đã có kết quả VÀ dữ liệu đã sửa sau lần chạy tạo ra
        nó (hoặc CSDL cũ chưa từng ghi dấu vân tay). Dùng cho dải cảnh báo ở
        thẻ Kết quả, cho mọi lần xuất tệp, và cho tệp tổng hợp.
        """
        try:
            with self._ket_noi_doc() as cur:
                co_ket_qua = cur.execute(
                    "SELECT COUNT(*) FROM match_results").fetchone()[0] > 0
                dvt = cur.execute(
                    "SELECT dau_van_tay, luc FROM dau_van_tay_chay WHERE id = 1"
                ).fetchone()
                meta = cur.execute(
                    "SELECT run_at FROM run_meta WHERE id = 1").fetchone()
                hien_tai = self._dau_van_tay_du_lieu(cur) if co_ket_qua else None
            ket_qua_cu = bool(co_ket_qua and (dvt is None or dvt[0] != hien_tai))
            return _ok({
                "co_ket_qua": co_ket_qua,
                "ket_qua_cu": ket_qua_cu,
                "luc_chay": meta[0] if meta else None,
            })
        except Exception as e:
            return _fail(err("error_reading_results", detail=str(e)))
    # Thứ tự cố định để tệp và màn hình liệt kê giống nhau.
    _LOAI_THAY_DOI = ("moi_duoc_xep", "mat_cho", "doi_clb", "doi_dien",
                      "hoc_sinh_moi", "bi_bo")
    def get_thay_doi_ket_qua(self):
        """So bảng kết quả hiện tại với bản chụp NGAY TRƯỚC lần chạy gần nhất.

        So theo cặp (học sinh, buổi). Chỉ trả những cặp KHÁC nhau:
          * `moi_duoc_xep` — trước không có CLB, nay có
          * `mat_cho`      — trước có CLB, nay không
          * `doi_clb`      — CLB khác
          * `doi_dien`     — cùng CLB, đổi diện (dự trữ ↔ thường)
          * `hoc_sinh_moi` — trước chưa có dòng nào của cặp này
          * `bi_bo`        — nay không còn dòng nào của cặp này
        Lần chạy đầu tiên thì bản chụp rỗng -> `co_lan_truoc` = False.
        """
        try:
            with self._ket_noi_doc() as cur:
                meta = cur.execute(
                    "SELECT run_at FROM run_meta WHERE id = 1").fetchone()
                truoc = {
                    (r["student_id"], r["buoi"]): r for r in cur.execute(
                        "SELECT * FROM ket_qua_truoc")
                }
                nay = {
                    (r["student_id"], r["buoi"]): r for r in cur.execute(
                        "SELECT * FROM match_results")
                }
                ten_hs = dict(cur.execute(
                    "SELECT student_id, name FROM students").fetchall())
                ten_clb = dict(cur.execute(
                    "SELECT club_id, name FROM clubs").fetchall())
                run_at_truoc = next(
                    (r["run_at"] for r in truoc.values() if r["run_at"]), None)
                ds_buoi = self._ds_buoi(cur)

            dong = []
            for khoa in sorted(set(truoc) | set(nay)):
                a, b = truoc.get(khoa), nay.get(khoa)
                clb_a = a["club_id"] if a else None
                clb_b = b["club_id"] if b else None
                dien_a = a["matched_tier"] if a else None
                dien_b = b["matched_tier"] if b else None
                if a is None:
                    loai = "hoc_sinh_moi"
                elif b is None:
                    loai = "bi_bo"
                elif clb_a == clb_b:
                    if clb_a is None or dien_a == dien_b:
                        continue
                    loai = "doi_dien"
                elif clb_a is None:
                    loai = "moi_duoc_xep"
                elif clb_b is None:
                    loai = "mat_cho"
                else:
                    loai = "doi_clb"
                sid, buoi = khoa
                dong.append({
                    "student_id": sid, "name": ten_hs.get(sid) or "",
                    "buoi": buoi,
                    "club_cu": clb_a, "ten_club_cu": ten_clb.get(clb_a) if clb_a else None,
                    "dien_cu": dien_a,
                    "club_moi": clb_b, "ten_club_moi": ten_clb.get(clb_b) if clb_b else None,
                    "dien_moi": dien_b,
                    "loai": loai,
                })
            return _ok({
                "co_lan_truoc": bool(truoc),
                "run_at_truoc": run_at_truoc,
                "run_at_moi": meta[0] if meta else None,
                "nhieu_buoi": len(ds_buoi) > 1,
                "dong": dong,
                "dem": {k: sum(1 for d in dong if d["loai"] == k)
                        for k in self._LOAI_THAY_DOI},
            })
        except Exception as e:
            return _fail(err("error_reading_results", detail=str(e)))
    _NHAN_LOAI_THAY_DOI = {
        "moi_duoc_xep": "Mới được xếp",
        "mat_cho": "Mất chỗ",
        "doi_clb": "Đổi CLB",
        "doi_dien": "Đổi diện",
        "hoc_sinh_moi": "Học sinh mới",
        "bi_bo": "Không còn trong kết quả",
    }
    @staticmethod
    def _nhan_dien(tier):
        # 'reserve'/'general' la ma noi bo — giao vien khong phai doan.
        return {"reserve": "Dự trữ", "general": "Thường"}.get(tier or "", "")
    def export_thay_doi_ket_qua(self, output_path: str = ""):
        """Tệp `thay_doi_ket_qua.csv`: những em có kết quả khác lần chạy trước."""
        try:
            kq = self.get_thay_doi_ket_qua()
            if not kq["ok"]:
                return kq
            d = kq["data"]
            if not d["co_lan_truoc"]:
                return _fail(err("chua_co_lan_chay_truoc"))
            path = self._duong_xuat(output_path, "thay_doi_ket_qua.csv")
            nb = d["nhieu_buoi"]
            header = (["Mã học sinh", "Họ tên"] + (["Buổi"] if nb else [])
                      + ["CLB cũ", "Diện cũ", "CLB mới", "Diện mới", "Loại thay đổi"])
            dong = [
                [x["student_id"], x["name"]] + ([x["buoi"]] if nb else [])
                + [x["ten_club_cu"] or x["club_cu"] or "(không có)",
                   self._nhan_dien(x["dien_cu"]),
                   x["ten_club_moi"] or x["club_moi"] or "(không có)",
                   self._nhan_dien(x["dien_moi"]),
                   self._NHAN_LOAI_THAY_DOI[x["loai"]]]
                for x in d["dong"]
            ]
            tt = self.get_trang_thai_ket_qua()
            ket_qua_cu = bool(tt["ok"] and tt["data"]["ket_qua_cu"])
            # Dong dau noi RO so voi lan nao — thieu no thi tep "thay doi"
            # khong noi duoc la thay doi tu dau.
            dau = [["So sánh kết quả lúc", d["run_at_truoc"] or "",
                    "với kết quả lúc", d["run_at_moi"] or ""]]
            if ket_qua_cu:
                dau.append(["CẢNH BÁO: dữ liệu đã sửa sau lần chạy gần nhất — "
                            "kết quả có thể đã cũ, hãy chạy lại."])
            _ghi_csv(path, [], dau + [[]] + [header] + dong)
            return _ok({"path": path, "n_thay_doi": len(dong),
                        "dem": d["dem"], "ket_qua_cu": ket_qua_cu})
        except Exception as e:
            return _fail(err("error_exporting_csv", detail=str(e)))

    # -----------------------------------------------------------------
    # DỮ LIỆU ĐẦU VÀO HIỆN TẠI — ba tệp NẠP LẠI ĐƯỢC
    # -----------------------------------------------------------------
    def export_du_lieu_dau_vao(self, output_path: str = ""):
        """Xuất dữ liệu đầu vào ĐANG CÓ (kể cả sửa sau lần chạy) thành ba tệp.

        Mỗi tệp đúng một định dạng mà cửa nạp đã hiểu (dạng dài, tiêu đề như
        `mau_csv/`), nên kéo ba tệp này vào ô nạp của một CSDL trống là dựng
        lại y nguyên — test `test_xuat_du_lieu.py` canh điều đó:

          01_danh_sach_CLB.csv       club_id,name,capacity,reserve_capacity,reserve_group,buoi
          02_chon_CLB_muon_thi.csv   student_id,name,reserve_group,test_club_1,score_1,...
          03_xep_hang_nguyen_vong.csv student_id,name,reserve_group,club_id,rank

        Em không có ô tick nào lẫn nguyện vọng nào thì KHÔNG định dạng nạp
        nào chứa được (mọi dòng cần một club_id) — đếm vào
        `n_hs_khong_co_dong` để giao diện nói ra; tệp hồ sơ học sinh vẫn có em.
        """
        try:
            ten = "du_lieu_dau_vao_" + datetime.datetime.now().strftime("%Y%m%d_%H%M")
            thu_muc = self._duong_xuat(output_path, ten)
            os.makedirs(thu_muc, exist_ok=True)
            with self._ket_noi_doc() as cur:
                clb = cur.execute(
                    "SELECT club_id, name, capacity, reserve_capacity, "
                    "COALESCE(reserve_group, ''), buoi FROM clubs ORDER BY club_id"
                ).fetchall()
                tick = cur.execute("""
                    SELECT t.student_id, s.name, COALESCE(s.reserve_group, ''),
                           t.club_id, sc.score
                    FROM club_test_selection t
                    JOIN students s ON s.student_id = t.student_id
                    LEFT JOIN club_scores sc
                           ON sc.student_id = t.student_id AND sc.club_id = t.club_id
                    ORDER BY t.student_id, t.club_id
                """).fetchall()
                nv = cur.execute("""
                    SELECT p.student_id, s.name, COALESCE(s.reserve_group, ''),
                           p.club_id, p.rank
                    FROM preferences p
                    JOIN students s ON s.student_id = p.student_id
                    ORDER BY p.student_id, p.rank
                """).fetchall()
                n_trong = cur.execute("""
                    SELECT COUNT(*) FROM students s
                    WHERE NOT EXISTS (SELECT 1 FROM club_test_selection t
                                      WHERE t.student_id = s.student_id)
                      AND NOT EXISTS (SELECT 1 FROM preferences p
                                      WHERE p.student_id = s.student_id)
                """).fetchone()[0]

            def so(x):
                # 8.0 -> "8", 8.5 -> "8.5": dung cach viet giao vien go vao.
                return "" if x is None else ("%g" % x)

            duong = {
                "clubs": os.path.join(thu_muc, "01_danh_sach_CLB.csv"),
                "test_selection": os.path.join(thu_muc, "02_chon_CLB_muon_thi.csv"),
                "preferences": os.path.join(thu_muc, "03_xep_hang_nguyen_vong.csv"),
            }
            _ghi_csv(duong["clubs"],
                     ["club_id", "name", "capacity", "reserve_capacity",
                      "reserve_group", "buoi"],
                     [[r[0], r[1] or "", r[2], r[3], r[4],
                       "" if (r[5] or BUOI_MAC_DINH) == BUOI_MAC_DINH else r[5]]
                      for r in clb])
            # DANG RONG (test_club_N, score_N), khong phai dang dai
            # (club_id, score): dang dai co y bi cua nap coi la MO HO va hoi
            # lai nguoi dung (tests/test_upload_tu_nhan_dien.py). Dang rong
            # thi nhan dien CHAC, keo vao la nap — dung mau
            # `mau_csv/vi_du_day_du/02_chon_CLB_muon_thi.csv`.
            theo_em: dict = {}
            for r in tick:
                em = theo_em.setdefault(r[0], [r[1] or "", r[2], []])
                em[2].append((r[3], so(r[4])))
            n_cot = max((len(e[2]) for e in theo_em.values()), default=1)
            _ghi_csv(duong["test_selection"],
                     ["student_id", "name", "reserve_group"]
                     + [c for i in range(1, n_cot + 1)
                        for c in ("test_club_%d" % i, "score_%d" % i)],
                     [[sid, e[0], e[1]]
                      + [o for cid_d in e[2] for o in cid_d]
                      + [""] * (2 * (n_cot - len(e[2])))
                      for sid, e in sorted(theo_em.items())])
            _ghi_csv(duong["preferences"],
                     ["student_id", "name", "reserve_group", "club_id", "rank"],
                     [[r[0], r[1] or "", r[2], r[3], r[4]] for r in nv])
            return _ok({
                "dir": thu_muc,
                "paths": duong,
                "n_clubs": len(clb),
                "n_tick": len(tick),
                "n_scores": sum(1 for r in tick if r[4] is not None),
                "n_preferences": len(nv),
                "n_hs_khong_co_dong": n_trong,
            })
        except Exception as e:
            return _fail(err("error_exporting_csv", detail=str(e)))

    # -----------------------------------------------------------------
    # HỒ SƠ HỌC SINH — một em hoặc nhiều em được chọn
    # -----------------------------------------------------------------
    def export_hoc_sinh_csv(self, student_ids: list, output_path: str = ""):
        """Toàn bộ dữ liệu của các em được chọn, mỗi dòng một cặp (em, CLB).

        Một cặp có mặt khi em tick thi, xếp nguyện vọng, HOẶC được xếp vào
        CLB đó. Em chưa có gì vẫn có MỘT dòng — chọn em nào thì tệp phải có
        em đó, thiếu im lặng là người nhận tưởng em không tồn tại.

        Điểm: chỉ điểm còn tick, đúng tập lần chạy dùng. Số bốc thăm: trường
        nhiều buổi dùng số TỪNG BUỔI (cùng nguồn với bảng số bốc thăm), một
        buổi dùng số đã khoá. Nguyện vọng thứ: thứ hạng TRONG BUỔI, như cột
        cùng tên ở bảng kết quả.
        """
        try:
            if not isinstance(student_ids, list) or not student_ids:
                return _fail(err("chua_chon_hoc_sinh"))
            ids = list(dict.fromkeys(str(s) for s in student_ids))

            with self._ket_noi_doc() as cur:
                hs = {r["student_id"]: r for r in cur.execute(
                    "SELECT student_id, name, reserve_group, stb_number FROM students "
                    "WHERE student_id IN (%s)" % ",".join("?" * len(ids)), ids)}
                thieu = [s for s in ids if s not in hs]
                if thieu:
                    return _fail(err("student_not_found", student_id=", ".join(thieu)))
                clb = {r["club_id"]: r for r in cur.execute(
                    "SELECT club_id, name, COALESCE(NULLIF(buoi, ''), ?) AS buoi "
                    "FROM clubs", (BUOI_MAC_DINH,))}
                ds_buoi = self._ds_buoi(cur)
                dau = ",".join("?" * len(ids))
                tick = {(r[0], r[1]) for r in cur.execute(
                    "SELECT student_id, club_id FROM club_test_selection "
                    "WHERE student_id IN (%s)" % dau, ids)}
                diem = {(r[0], r[1]): r[2] for r in cur.execute(
                    "SELECT student_id, club_id, score FROM club_scores "
                    "WHERE student_id IN (%s)" % dau, ids)}
                nv: dict = {}
                for r in cur.execute(
                        "SELECT student_id, club_id FROM preferences "
                        "WHERE student_id IN (%s) ORDER BY student_id, rank" % dau, ids):
                    nv.setdefault(r[0], []).append(r[1])
                kq = {(r["student_id"], r["club_id"]): r for r in cur.execute(
                    "SELECT student_id, club_id, matched_tier FROM match_results "
                    "WHERE club_id IS NOT NULL AND student_id IN (%s)" % dau, ids)}

            nhieu_buoi = len(ds_buoi) > 1
            so_buoi: dict = {}
            if nhieu_buoi:
                tham = self.get_so_boc_tham_theo_buoi()
                if tham["ok"]:
                    so_buoi = {e["student_id"]: e["so"]
                               for e in tham["data"]["hoc_sinh"]}

            def buoi_cua(cid):
                return clb[cid]["buoi"] if cid in clb else BUOI_MAC_DINH

            header = (["Mã học sinh", "Họ tên", "Nhóm dự trữ"]
                      + (["Buổi"] if nhieu_buoi else [])
                      + ["Số bốc thăm", "Mã CLB", "Tên CLB", "Đăng ký thi",
                         "Điểm", "Nguyện vọng thứ", "Kết quả", "Diện"])
            dong = []
            for sid in ids:
                h = hs[sid]
                # Thu hang trong buoi: dem theo thu tu nguyen vong cua em.
                hang: dict = {}
                dem_buoi: dict = {}
                for cid in nv.get(sid, []):
                    b = buoi_cua(cid)
                    dem_buoi[b] = dem_buoi.get(b, 0) + 1
                    hang[cid] = dem_buoi[b]
                cac_clb = list(dict.fromkeys(
                    nv.get(sid, [])
                    + sorted(c for (s, c) in tick if s == sid)
                    + sorted(c for (s, c) in kq if s == sid)))
                # Sap theo buoi (thu tu ngay), roi theo thu tu tren.
                cac_clb.sort(key=lambda c: (ds_buoi.index(buoi_cua(c))
                                            if buoi_cua(c) in ds_buoi else 99))

                def so_tham(b, sid=sid, h=h):
                    if nhieu_buoi:
                        v = so_buoi.get(sid, {}).get(b)
                        return "" if v is None else v
                    return "" if h["stb_number"] is None else h["stb_number"]

                goc = [sid, h["name"] or "", h["reserve_group"] or ""]
                if not cac_clb:
                    dong.append(goc + ([""] if nhieu_buoi else [])
                                + [so_tham(None) if not nhieu_buoi else "",
                                   "", "(chưa nhập gì)", "", "", "", "", ""])
                    continue
                for cid in cac_clb:
                    b = buoi_cua(cid)
                    co_tick = (sid, cid) in tick
                    d = diem.get((sid, cid)) if co_tick else None
                    x = kq.get((sid, cid))
                    dong.append(
                        goc + ([b] if nhieu_buoi else [])
                        + [so_tham(b), cid,
                           clb[cid]["name"] if cid in clb else "",
                           "Có" if co_tick else "—",
                           "" if d is None else ("%g" % d),
                           hang.get(cid, ""),
                           "Được xếp" if x else "—",
                           self._nhan_dien(x["matched_tier"]) if x else ""])

            ten = ("hoc_sinh_%s.csv" % self._ten_file_an_toan(ids[0], "hoc_sinh")
                   if len(ids) == 1 else "hoc_sinh_da_chon_%d_em.csv" % len(ids))
            path = self._duong_xuat(output_path, ten)
            _ghi_csv(path, header, dong)
            tt = self.get_trang_thai_ket_qua()
            return _ok({
                "path": path,
                "n_hoc_sinh": len(ids),
                "n_dong": len(dong),
                "ket_qua_cu": bool(tt["ok"] and tt["data"]["ket_qua_cu"]),
            })
        except Exception as e:
            return _fail(err("error_exporting_csv", detail=str(e)))

    def export_csv(self, output_path: str = ""):
        """Giữ khoá ghi + bật bộ nhớ báo cáo suốt lần xuất.

        Khoá ghi: không lời gọi ghi nào chen vào giữa, nên mọi tệp của một
        lần xuất cùng chụp MỘT trạng thái dữ liệu. Bộ nhớ: xem
        _nho_trong_luc_xuat.
        """
        with self._khoa_ghi:
            self._bo_nho_xuat = {}
            try:
                return self._export_csv_da_khoa(output_path)
            finally:
                self._bo_nho_xuat = None

    def _export_csv_da_khoa(self, output_path: str = ""):
        """
        Xuất kết quả phân bổ ra CSV cho nhà trường:

          1. MỘT FILE TỔNG — mọi học sinh, đủ tên và diện trúng tuyển.
          2. MỘT THƯ MỤC theo CLB — mỗi CLB một file để phát cho giáo
             viên phụ trách, kèm file `_chua_duoc_xep.csv` liệt kê các
             em chưa vào CLB nào (nhóm nhà trường cần xử lý tiếp).

        output_path rỗng hoặc TƯƠNG ĐỐI -> tệp rơi vào **thư mục Tải
        xuống** của người dùng. Đường dẫn tương đối vốn rơi vào thư mục
        làm việc của tiến trình — trên máy Windows chạy .exe qua
        shortcut, thư mục đó có thể là bất kỳ đâu, người dùng không tìm
        ra tệp.

        Trước đây tệp được đặt CẠNH app.db. Chỗ đó tìm được, nhưng nằm
        trong thư mục cài đặt phần mềm — không phải chỗ để tệp cho người
        ta mang đi, và lẫn với `.exe` cùng dữ liệu. Thư mục Tải xuống là
        chỗ ai cũng biết tìm.

        Không tìm được thư mục Tải xuống (hoặc không ghi được vào đó) thì
        LÙI VỀ cạnh app.db — không bao giờ để việc xuất kết quả thất bại
        chỉ vì chuyện chỗ để tệp.

        Đường dẫn TUYỆT ĐỐI do bên gọi truyền vào thì tôn trọng nguyên
        văn, kể cả việc ghi đè. Luôn trả về đường dẫn ĐẦY ĐỦ để giao
        diện hiện đúng chỗ tệp nằm.
        """
        try:
            import csv

            output_path = self._duong_xuat(output_path, "ket_qua_phan_bo.csv")

            with self._ket_noi_doc() as cur:
                rows = cur.execute("""
                    SELECT m.student_id, s.name AS ho_ten, m.buoi, m.club_id,
                           c.name AS ten_club, m.rank_in_student_pref, m.matched_tier,
                           s.reserve_group
                    FROM match_results m
                    LEFT JOIN students s ON s.student_id = m.student_id
                    LEFT JOIN clubs   c ON c.club_id   = m.club_id
                    ORDER BY m.student_id, m.buoi COLLATE THU_TU_BUOI
                """).fetchall()
                ds_buoi = self._ds_buoi(cur)
            nhieu_buoi = len(ds_buoi) > 1

            an_toan_cho_excel = _an_toan_cho_excel

            def dien(tier):
                # 'reserve'/'general' la ma noi bo — giao vien khong phai doan.
                return {"reserve": "Dự trữ", "general": "Thường"}.get(tier or "", "")

            # Cot "Buoi" chi xuat hien khi truong THAT SU dung nhieu buoi.
            # Mot buoi ma van co cot day gia tri "__mac_dinh__" thi giao vien
            # phai doan xem no nghia la gi — them mot cot vo nghia vao tep ai
            # cung phai doc la mot cai gia that.
            COT_TONG = (["Mã học sinh", "Họ tên"]
                        + (["Buổi"] if nhieu_buoi else [])
                        + ["Mã CLB", "Tên CLB", "Nguyện vọng thứ",
                           "Diện trúng tuyển", "Nhóm dự trữ"])
            COT_CLB = ["Mã học sinh", "Họ tên", "Nguyện vọng thứ",
                       "Diện trúng tuyển", "Nhóm dự trữ"]

            def dong_tong(r):
                return (
                    [r["student_id"], r["ho_ten"] or ""]
                    + ([r["buoi"]] if nhieu_buoi else [])
                    + [
                        r["club_id"] or "",
                        r["ten_club"] or ("" if r["club_id"] else "(chưa được xếp)"),
                        r["rank_in_student_pref"] if r["rank_in_student_pref"] else "",
                        dien(r["matched_tier"]), r["reserve_group"] or "",
                    ]
                )

            ghi = _ghi_csv   # UTF-8 co BOM + chan cong thuc Excel

            # Giu lai tung bang de con dung lai cho tep Excel o cuoi ham —
            # dung MOT nguon du lieu cho ca hai dinh dang, khong dung hai
            # duong tinh song song roi troi khoi nhau.
            bang_xuat: list[tuple] = []
            dong_tong_tat_ca = [dong_tong(r) for r in rows]
            ghi(output_path, COT_TONG, dong_tong_tat_ca)
            bang_xuat.append(("Bảng tổng", COT_TONG, dong_tong_tat_ca))

            goc, _ = os.path.splitext(output_path)
            per_club_dir = goc + "_theo_club"
            os.makedirs(per_club_dir, exist_ok=True)

            # Xoa tep .csv cua lan xuat TRUOC trong dung thu muc nay.
            # Khong xoa thi mot CLB da bi go khoi he thong van con nguyen
            # tep cua no, nam canh cac tep moi va trong y het nhu that.
            # Giao vien cam nham tep do di to chuc mot CLB khong con ton
            # tai — sai ma khong co dau hieu nao bao.
            # Chi dung .csv, chi trong thu muc do phan mem tu tao ra.
            for cu_ten in os.listdir(per_club_dir):
                if cu_ten.lower().endswith(".csv"):
                    try:
                        os.remove(os.path.join(per_club_dir, cu_ten))
                    except OSError:
                        pass

            theo_club: dict = {}
            for r in rows:
                theo_club.setdefault(r["club_id"] or None, []).append(r)

            n_file = 0
            # Giu cho "_chua_duoc_xep" truoc, de mot CLB trung ten do khong
            # ghi de danh sach em chua co cho.
            ten_da_dung: set = {"_chua_duoc_xep"}
            for club_id, ds in theo_club.items():
                if club_id is None:
                    ten = "_chua_duoc_xep"
                else:
                    # Hai club_id khac nhau co the ra cung ten sau khi lam
                    # sach ("a/b" va "a\b"), hoac chi khac hoa/thuong.
                    ten = self._ten_chua_dung(
                        self._ten_file_an_toan(club_id), ten_da_dung)
                ghi(
                    os.path.join(per_club_dir, ten + ".csv"),
                    COT_CLB,
                    [[r["student_id"], r["ho_ten"] or "",
                      r["rank_in_student_pref"] if r["rank_in_student_pref"] else "",
                      dien(r["matched_tier"]), r["reserve_group"] or ""] for r in ds],
                )
                n_file += 1

            # ---- Thoi khoa bieu tuan: SAN PHAM CHINH cua ban nhieu buoi ----
            # Moi em MOT dong, moi buoi MOT cot — dan bang duoc, phat cho hoc
            # sinh duoc. Bang tong o tren van giu nguyen hinh dang cu (moi
            # dong la mot cap em-buoi) vi do la thu de doi chieu va loc, con
            # day la thu de doc.
            duong_tkb = None
            n_tkb = 0
            duong_tham = None
            n_tham = 0
            if nhieu_buoi:
                theo_em: dict = {}
                for r in rows:
                    em = theo_em.setdefault(
                        r["student_id"], {"ho_ten": r["ho_ten"] or "", "o": {}})
                    em["o"][r["buoi"]] = (
                        r["ten_club"] or r["club_id"] or "") if r["club_id"] else ""
                duong_tkb = goc + "_thoi_khoa_bieu.csv"
                COT_TKB = ["Mã học sinh", "Họ tên"] + list(ds_buoi) + ["Số CLB"]
                DONG_TKB = [
                    [sid, em["ho_ten"]]
                    + [em["o"].get(b, "") for b in ds_buoi]
                    + [sum(1 for b in ds_buoi if em["o"].get(b))]
                    for sid, em in sorted(theo_em.items())
                ]
                ghi(duong_tkb, COT_TKB, DONG_TKB)
                bang_xuat.append(("Thời khoá biểu", COT_TKB, DONG_TKB))
                n_tkb = len(theo_em)

                # ---- Mot tep moi buoi, cho giao vien truc ngay do ----
                per_buoi_dir = goc + "_theo_buoi"
                os.makedirs(per_buoi_dir, exist_ok=True)
                for cu_ten in os.listdir(per_buoi_dir):
                    if cu_ten.lower().endswith(".csv"):
                        try:
                            os.remove(os.path.join(per_buoi_dir, cu_ten))
                        except OSError:
                            pass
                for b in ds_buoi:
                    ghi(
                        os.path.join(per_buoi_dir, self._ten_file_an_toan(b, "buoi") + ".csv"),
                        ["Mã học sinh", "Họ tên", "Mã CLB", "Tên CLB",
                         "Nguyện vọng thứ", "Diện trúng tuyển"],
                        [
                            [r["student_id"], r["ho_ten"] or "", r["club_id"] or "",
                             r["ten_club"] or "",
                             r["rank_in_student_pref"] if r["rank_in_student_pref"] else "",
                             dien(r["matched_tier"])]
                            for r in rows
                            if r["buoi"] == b and r["club_id"]
                        ],
                    )

                # ---- So boc tham tung buoi: TEP DE TRUONG DAN BANG ----
                # Phan mem xao lai thu tu o moi buoi, nen se co phu huynh hoi
                # "vi sao con toi thu Ba dung thu 30 ma thu Sau dung thu 120".
                # Khong dua duoc tep nay ra thi cach boc tham do khong dung
                # duoc o truong, du phep do noi no cong bang hon.
                #
                # Truong van chi cong bo MOT thu: bo so da khoa cong hat giong.
                # Bang nay la ham TAT DINH cua hai thu do, ai cung tinh lai duoc.
                kq_tham = self.get_so_boc_tham_theo_buoi()
                if kq_tham["ok"] and kq_tham["data"]["da_chay"]:
                    duong_tham = goc + "_so_boc_tham_theo_buoi.csv"
                    COT_THAM = ["Mã học sinh", "Họ tên"] + list(ds_buoi)
                    DONG_THAM = [
                        [e["student_id"], e["name"] or ""]
                        + [e["so"].get(b, "") for b in ds_buoi]
                        for e in kq_tham["data"]["hoc_sinh"]
                    ]
                    bang_xuat.append(("Số bốc thăm", COT_THAM, DONG_THAM))
                    ghi(
                        duong_tham,
                        COT_THAM,
                        [
                            [e["student_id"], e["name"] or ""]
                            + [e["so"].get(b, "") for b in ds_buoi]
                            for e in kq_tham["data"]["hoc_sinh"]
                        ],
                    )
                    n_tham = len(kq_tham["data"]["hoc_sinh"])

            # ---- Tep TONG HOP: ket qua nay den tu dau, va con gi phai lam ----
            # Xuat cho MOI truong, mot buoi hay nhieu buoi deu can. Xem
            # `_bang_tong_hop` de biet vi sao.
            duong_tong_hop = goc + "_tong_hop.csv"
            bang = self._bang_tong_hop()
            with open(duong_tong_hop, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerows([an_toan_cho_excel(o) for o in d] for d in bang)

            # ---- Bon trang Excel cho bon muc moi cua the Ket qua ----
            # Man hinh va tep phai noi CUNG mot con so: ca hai doc tu cung
            # bon ham, khong co duong nao de chung troi khoi nhau.
            ten_trang = {"clb": "Theo câu lạc bộ", "nv": "Nguyện vọng",
                         "dt": "Suất dự trữ", "tt": "Chưa có chỗ"}
            for khoa, (_, cot, dong) in self._bang_bao_cao(so=True).items():
                if dong:
                    bang_xuat.append((ten_trang[khoa], cot, dong))

            # ---- MOT tep Excel gom moi bang ca truong ----
            # Xuat hong KHONG duoc lam hong ca lan xuat: cac tep .csv da ghi
            # xong va van dung. Bao lai bang duong dan None chu khong nem loi.
            duong_excel = goc + ".xlsx"
            try:
                self._xuat_excel(
                    duong_excel,
                    # Tong hop dung dau: mo tep ra la thay ngay ket qua den tu
                    # dau. Bang tong hop khong co dong tieu de rieng (no tu
                    # chia phan), nen truyen cot rong.
                    [("Tổng hợp", [], bang)] + bang_xuat,
                )
            except Exception:
                duong_excel = None

            return _ok({
                "path": output_path,
                "tong_hop_path": duong_tong_hop,
                "excel_path": duong_excel,
                "n_rows": len(rows),
                "per_club_dir": per_club_dir,
                "n_club_files": n_file,
                "nhieu_buoi": nhieu_buoi,
                "thoi_khoa_bieu_path": duong_tkb,
                "n_thoi_khoa_bieu_rows": n_tkb,
                "so_boc_tham_path": duong_tham,
                "n_so_boc_tham_rows": n_tham,
            })
        except Exception as e:
            return _fail(err("error_exporting_csv", detail=str(e)))
