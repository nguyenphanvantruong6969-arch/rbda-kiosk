"""Sinh KHUNG poster ba lá (trống nội dung) cho gian trưng bày.

Kích thước thật khi mở phẳng: 200 cm x 150 cm
    lá trái 50 cm | lá giữa 100 cm | lá phải 50 cm
Gập hai lá bên vào 90 độ: rộng 100 cm, sâu 50 cm, cao 150 cm.

Tệp này CHỈ dựng khung: lưới, màu, cỡ chữ, ô chờ. Không có nội dung đề tài —
Phụ lục 1 cấm AI viết bản thảo poster; học sinh tự viết chữ vào các ô.

Chạy:
    python poster/tao_poster.py

Sinh ra trong poster/:
    poster_khung.html          trang in, đúng tỉ lệ 1:1 (đơn vị cm)
    poster_khung_200x150cm.pdf bản in tỉ lệ 1:1, vector
    poster_huong_dan.pdf       cùng khung, phủ nếp gấp, lề an toàn, kích thước
    poster_khung_1-2.pptx      bản sửa được, tỉ lệ 1:2 (100 x 75 cm) — in phóng 200%
    xem_truoc_phang.png        ảnh xem trước khi mở phẳng
    xem_truoc_gap.png          ảnh minh hoạ khi đã gập
"""

from pathlib import Path

HERE = Path(__file__).resolve().parent

# ---------------------------------------------------------------- kích thước
W, H = 200.0, 150.0          # cm, mở phẳng
FOLDS = (50.0, 150.0)        # vị trí hai nếp gấp
MARGIN = 3.0                 # lề ngoài
FOLD_GAP = 2.5               # khoảng trống mỗi bên nếp gấp

# Cỡ chữ ở tỉ lệ 1:1 (pt). Bản pptx 1:2 lấy một nửa.
PT_TITLE = 140
PT_KICKER = 40
PT_SUB = 44
PT_HEAD = 60
PT_BODY = 32
PT_SMALL = 26

# ---------------------------------------------------------------- bảng màu
# Lấy đúng token của giao diện phần mềm (style.css): giấy lạnh + ink + vàng đồng
INK = "#1C2321"
INK_SOFT = "#3A423F"
SLATE = "#5B6570"
PAPER = "#F3F4F1"
RAISED = "#FFFFFF"
HAIRLINE = "#DBDCD5"
GOLD = "#C98A1F"
GOLD_SOFT = "#F1DFB8"
MOSS = "#3F6B52"
MOSS_SOFT = "#DCE9E0"
RUST = "#B84A3E"

# ---------------------------------------------------------------- các ô
# x, y, w, h tính bằng cm trên mặt phẳng 200 x 150.
# hint_h: chiều cao vùng gợi ý; phần còn lại là ô ảnh (nếu có image).
LX, LW = MARGIN, FOLDS[0] - FOLD_GAP - MARGIN                    # 3 → 47,5
CX, CW = FOLDS[0] + FOLD_GAP, FOLDS[1] - FOLDS[0] - 2 * FOLD_GAP  # 52,5 → 147,5
RX, RW = FOLDS[1] + FOLD_GAP, W - MARGIN - FOLDS[1] - FOLD_GAP    # 152,5 → 197
COL_GAP = 2.5
CW2 = (CW - COL_GAP) / 2

HEADER = dict(x=CX, y=MARGIN, w=CW, h=27.0)

BLOCKS = [
    # ---- lá trái
    dict(num="01", title="Đặt vấn đề", x=LX, y=3.0, w=LW, h=44.0,
         hint="Vấn đề thực tế là gì? Ai gặp phải, ở đâu, mức độ ra sao? "
              "Vì sao cách làm hiện tại chưa ổn?"),
    dict(num="02", title="Mục tiêu & câu hỏi nghiên cứu", x=LX, y=49.5, w=LW, h=38.0,
         hint="Mục tiêu chính và 2–3 câu hỏi nghiên cứu, mỗi câu một dòng."),
    dict(num="03", title="Cơ sở lý thuyết", x=LX, y=90.0, w=LW, h=56.0,
         hint="Khái niệm và tài liệu nền tảng đề tài dựa vào.",
         hint_h=11.0, image="Sơ đồ khái niệm (nếu có)"),
    # ---- lá giữa
    dict(num="04", title="Phương pháp & quy trình", x=CX, y=32.5, w=CW, h=42.0,
         hint="Các bước thực hiện, từ dữ liệu đầu vào tới kết quả.",
         hint_h=8.5, image="Sơ đồ quy trình — khổ ngang"),
    dict(num="05", title="Sản phẩm", x=CX, y=77.0, w=CW2, h=69.0,
         hint="Sản phẩm làm được gì, dùng thế nào.",
         hint_h=9.0, image="Ảnh chụp giao diện / sản phẩm", image_h=38.0),
    dict(num="06", title="Kết quả", x=CX + CW2 + COL_GAP, y=77.0, w=CW2, h=69.0,
         hint="Số đo chính. Dữ liệu mô phỏng thì ghi rõ là mô phỏng.",
         hint_h=9.0, image="Biểu đồ / bảng số liệu", image_h=38.0),
    # ---- lá phải
    dict(num="07", title="Thảo luận & đánh giá", x=RX, y=3.0, w=RW, h=40.0,
         hint="Kết quả nói lên điều gì? Hạn chế của đề tài."),
    dict(num="08", title="Kết luận", x=RX, y=45.5, w=RW, h=40.0, accent=True,
         hint="Trả lời lần lượt từng câu hỏi nghiên cứu ở mục 02."),
    dict(num="09", title="Hướng phát triển", x=RX, y=88.0, w=RW, h=24.0,
         hint="Bước tiếp theo nếu tiếp tục đề tài."),
    dict(num="10", title="Tài liệu tham khảo", x=RX, y=114.5, w=RW, h=17.0,
         hint="Theo một chuẩn trích dẫn thống nhất.", small=True),
    dict(num="", title="Tuyên bố sử dụng AI", x=RX, y=134.0, w=RW, h=12.0,
         hint="Phần nào có AI hỗ trợ, theo Phụ lục 1.", small=True, muted=True),
]

BANNED = ["Logo nhà trường", "Tên trường", "Mã dự án",
          "Tên giáo viên hướng dẫn", "Tên học sinh tham gia"]


# ================================================================ HTML
def _pos(b):
    return (f"left:{b['x']}cm;top:{b['y']}cm;"
            f"width:{b['w']}cm;height:{b['h']}cm;")


def block_html(b):
    cls = "card"
    if b.get("accent"):
        cls += " accent"
    if b.get("muted"):
        cls += " muted"
    if b.get("small"):
        cls += " small"
    num = f'<span class="num">{b["num"]}</span>' if b["num"] else ""
    parts = [f'<section class="{cls}" style="{_pos(b)}">',
             f'<h2>{num}<span>{b["title"]}</span></h2>']
    hint_style = f' style="height:{b["hint_h"]}cm"' if b.get("image") else ""
    parts.append(f'<p class="hint"{hint_style}>{b["hint"]}</p>')
    if b.get("image"):
        img_style = f' style="height:{b["image_h"]}cm;flex:none"' if b.get("image_h") else ""
        parts.append(f'<div class="slot"{img_style}><span>[ {b["image"]} ]</span></div>')
        if b.get("image_h"):
            parts.append('<p class="hint after">Chú thích ảnh và 3–4 dòng diễn giải.</p>')
    parts.append("</section>")
    return "\n".join(parts)


def header_html():
    h = HEADER
    return f"""
<header class="head" style="{_pos(h)}">
  <p class="kicker">[ LĨNH VỰC DỰ THI ]</p>
  <h1>[ TÊN ĐỀ TÀI — VIẾT HOA,<br>TỐI ĐA HAI DÒNG ]</h1>
  <p class="sub">[ Một câu nói sản phẩm giải quyết việc gì ]</p>
</header>"""


def guide_html():
    """Lớp phủ hướng dẫn — chỉ có trong poster_huong_dan.pdf."""
    f1, f2 = FOLDS
    out = ['<div class="guide">']
    # vùng lề an toàn (tô chéo)
    out.append(f'<div class="g-safe" style="left:0;top:0;width:{W}cm;height:{MARGIN}cm"></div>')
    out.append(f'<div class="g-safe" style="left:0;top:{H - MARGIN}cm;width:{W}cm;height:{MARGIN}cm"></div>')
    out.append(f'<div class="g-safe" style="left:0;top:0;width:{MARGIN}cm;height:{H}cm"></div>')
    out.append(f'<div class="g-safe" style="left:{W - MARGIN}cm;top:0;width:{MARGIN}cm;height:{H}cm"></div>')
    for f in FOLDS:
        out.append(f'<div class="g-safe" style="left:{f - FOLD_GAP}cm;top:0;width:{2 * FOLD_GAP}cm;height:{H}cm"></div>')
        out.append(f'<div class="g-fold" style="left:{f}cm"></div>')
        out.append(f'<div class="g-foldlabel" style="left:{f}cm">NẾP GẤP</div>')
    # thước đo từng lá
    for x0, x1, label in [(0, f1, "50 cm"), (f1, f2, "100 cm"), (f2, W, "50 cm")]:
        out.append(f'<div class="g-dim" style="left:{x0}cm;width:{x1 - x0}cm"><span>{label}</span></div>')
    out.append(f'<div class="g-vdim"><span>150 cm</span></div>')
    # ghi chú
    banned = "".join(f"<li>{t}</li>" for t in BANNED)
    out.append(f"""
<div class="g-note" style="left:{f1 + 8}cm;top:{55}cm;width:{84}cm">
  <h3>Hướng dẫn in — lớp này KHÔNG in</h3>
  <ul class="cols">
    <li>Mở phẳng 200 × 150 cm. Gập hai lá bên vào 90° → rộng 100, sâu 50, cao 150 cm.</li>
    <li>Vùng gạch chéo: lề ngoài {MARGIN:g} cm và {FOLD_GAP:g} cm mỗi bên nếp gấp — không đặt chữ.</li>
    <li>Cỡ chữ tỉ lệ 1:1: tiêu đề {PT_TITLE} pt · mục {PT_HEAD} pt · nội dung {PT_BODY} pt · chú thích {PT_SMALL} pt.</li>
    <li>Ảnh chèn vào cần tối thiểu 100 ppi ở kích thước thật (ô 45 cm → ~1 800 px).</li>
  </ul>
  <h3 class="ban">Không được xuất hiện trên poster</h3>
  <ul class="ban">{banned}</ul>
</div>""")
    out.append("</div>")
    return "\n".join(out)


def build_html(guide=False):
    blocks = "\n".join(block_html(b) for b in BLOCKS)
    g = guide_html() if guide else ""
    return f"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8">
<title>Khung poster ba lá</title>
<link rel="stylesheet" href="../assets/fonts/fonts.css">
<style>
@page {{ size: {W}cm {H}cm; margin: 0; }}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
html, body {{ width:{W}cm; height:{H}cm; }}
body {{ background:{PAPER}; color:{INK}; font-family:"IBM Plex Sans",sans-serif;
       -webkit-print-color-adjust:exact; print-color-adjust:exact; }}
.board {{ position:relative; width:{W}cm; height:{H}cm; overflow:hidden; background:{PAPER}; }}

/* dải màu đầu mỗi lá — nhận diện ba lá khi đã gập */
.strip {{ position:absolute; top:0; height:1.2cm; }}

.card {{ position:absolute; background:{RAISED}; border:0.08cm solid {HAIRLINE};
        border-radius:0.6cm; padding:1.6cm 1.8cm; display:flex; flex-direction:column; gap:1cm; }}
.card h2 {{ font-family:"Be Vietnam Pro",sans-serif; font-weight:700; font-size:{PT_HEAD}pt;
           line-height:1.1; display:flex; align-items:center; gap:0.8cm; color:{INK}; }}
.card .num {{ flex:none; display:inline-grid; place-items:center; width:2.9cm; height:2.9cm;
             border-radius:50%; background:{GOLD}; color:{RAISED}; font-size:{PT_HEAD * 0.62:.0f}pt; }}
.card .hint {{ font-size:{PT_BODY}pt; line-height:1.4; color:{SLATE}; font-style:italic; flex:none; }}
.card .hint.after {{ font-size:{PT_SMALL}pt; }}
.card .slot {{ flex:1; border:0.12cm dashed {SLATE}; border-radius:0.4cm; background:{PAPER};
              display:grid; place-items:center; text-align:center; padding:1cm;
              font-size:{PT_SMALL}pt; color:{SLATE}; letter-spacing:0.02em; }}
.card.accent {{ background:{GOLD_SOFT}; border-color:{GOLD}; border-width:0.14cm; }}
.card.accent .num {{ background:{INK}; }}
.card.small {{ padding:1.2cm 1.8cm; gap:0.6cm; }}
.card.small h2 {{ font-size:{PT_HEAD * 0.7:.0f}pt; }}
.card.small .num {{ width:2.2cm; height:2.2cm; font-size:{PT_HEAD * 0.45:.0f}pt; }}
.card.small .hint {{ font-size:{PT_SMALL}pt; }}
.card.muted {{ background:{MOSS_SOFT}; border-color:{MOSS_SOFT}; }}
.card.muted h2 {{ color:{MOSS}; }}

.head {{ position:absolute; background:{INK}; color:{PAPER}; border-radius:0.6cm;
        padding:2.2cm 3cm; display:flex; flex-direction:column; justify-content:center; gap:0.9cm; }}
.head .kicker {{ font-family:"IBM Plex Sans",sans-serif; font-weight:600; font-size:{PT_KICKER}pt;
                letter-spacing:0.12em; color:{GOLD}; }}
.head h1 {{ font-family:"Be Vietnam Pro",sans-serif; font-weight:700; font-size:{PT_TITLE}pt;
           line-height:1.02; letter-spacing:-0.01em; }}
.head .sub {{ font-size:{PT_SUB}pt; color:{HAIRLINE}; }}

/* ------- lớp hướng dẫn ------- */
.guide {{ position:absolute; inset:0; pointer-events:none; }}
.g-safe {{ position:absolute; background:repeating-linear-gradient(45deg, rgba(184,74,62,.16) 0 0.5cm, transparent 0.5cm 1.2cm); }}
.g-fold {{ position:absolute; top:0; height:{H}cm; border-left:0.2cm dashed {RUST}; transform:translateX(-0.1cm); }}
.g-foldlabel {{ position:absolute; top:{H / 2}cm; transform:translate(-50%,-50%) rotate(-90deg);
               background:{RUST}; color:#fff; font:600 30pt "IBM Plex Sans"; padding:0.3cm 1.2cm; border-radius:0.3cm; }}
.g-dim {{ position:absolute; top:{H - 2.4}cm; height:1.8cm; border-left:0.12cm solid {RUST};
         border-right:0.12cm solid {RUST}; display:grid; place-items:center; }}
.g-dim::before {{ content:""; position:absolute; left:0; right:0; top:50%; border-top:0.12cm solid {RUST}; }}
.g-dim span, .g-vdim span {{ position:relative; background:{RUST}; color:#fff; font:600 34pt "IBM Plex Sans";
                            padding:0.1cm 1cm; border-radius:0.3cm; }}
.g-vdim {{ position:absolute; left:0.6cm; top:0; width:1.8cm; height:{H}cm; border-top:0.12cm solid {RUST};
          border-bottom:0.12cm solid {RUST}; display:grid; place-items:center; }}
.g-vdim::before {{ content:""; position:absolute; top:0; bottom:0; left:50%; border-left:0.12cm solid {RUST}; }}
.g-vdim span {{ transform:rotate(-90deg); white-space:nowrap; }}
.g-note {{ position:absolute; background:rgba(255,255,255,.96); border:0.2cm solid {RUST}; border-radius:0.6cm;
          padding:2cm 2.4cm; font-size:30pt; line-height:1.4; color:{INK}; }}
.g-note h3 {{ font:700 40pt "Be Vietnam Pro"; color:{RUST}; margin-bottom:0.8cm; }}
.g-note h3.ban {{ margin-top:1.4cm; }}
.g-note ul {{ padding-left:1.4cm; }}
.g-note ul.ban {{ display:flex; flex-wrap:wrap; gap:0.4cm 2.4cm; list-style:"✕  "; font-weight:600; }}
</style></head>
<body><div class="board">
<div class="strip" style="left:0;width:{FOLDS[0]}cm;background:{MOSS}"></div>
<div class="strip" style="left:{FOLDS[0]}cm;width:{FOLDS[1] - FOLDS[0]}cm;background:{GOLD}"></div>
<div class="strip" style="left:{FOLDS[1]}cm;width:{W - FOLDS[1]}cm;background:{MOSS}"></div>
{header_html()}
{blocks}
{g}
</div></body></html>"""


# ================================================================ PPTX 1:2
def build_pptx(path, s=0.5):
    from pptx import Presentation
    from pptx.util import Cm, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(W * s), Cm(H * s)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sh = slide.shapes

    def rect(x, y, w, h, fill, line=None, lw=0.08, rounded=True, dash=False):
        shape = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Cm(x * s), Cm(y * s), Cm(w * s), Cm(h * s))
        if rounded:
            shape.adjustments[0] = min(0.6 / min(w, h), 0.5)
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Cm(lw * s)
            if dash:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def text(x, y, w, h, runs, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
        tb = sh.add_textbox(Cm(x * s), Cm(y * s), Cm(w * s), Cm(h * s))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        for i, (t, pt, color, font, bold, italic) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = t
            r.font.size = Pt(pt * s)
            r.font.color.rgb = rgb(color)
            r.font.name = font
            r.font.bold = bold
            r.font.italic = italic
        return tb

    rect(0, 0, W, H, PAPER, rounded=False)
    rect(0, 0, FOLDS[0], 1.2, MOSS, rounded=False)
    rect(FOLDS[0], 0, FOLDS[1] - FOLDS[0], 1.2, GOLD, rounded=False)
    rect(FOLDS[1], 0, W - FOLDS[1], 1.2, MOSS, rounded=False)

    h = HEADER
    rect(h["x"], h["y"], h["w"], h["h"], INK)
    text(h["x"] + 3, h["y"] + 2.2, h["w"] - 6, h["h"] - 4.4, [
        ("[ LĨNH VỰC DỰ THI ]", PT_KICKER, GOLD, "IBM Plex Sans", True, False),
        ("[ TÊN ĐỀ TÀI — VIẾT HOA, TỐI ĐA HAI DÒNG ]", PT_TITLE, PAPER, "Be Vietnam Pro", True, False),
        ("[ Một câu nói sản phẩm giải quyết việc gì ]", PT_SUB, HAIRLINE, "IBM Plex Sans", False, False),
    ], anchor=MSO_ANCHOR.MIDDLE)

    for b in BLOCKS:
        small = b.get("small")
        fill = GOLD_SOFT if b.get("accent") else MOSS_SOFT if b.get("muted") else RAISED
        line = GOLD if b.get("accent") else MOSS_SOFT if b.get("muted") else HAIRLINE
        rect(b["x"], b["y"], b["w"], b["h"], fill, line, 0.14 if b.get("accent") else 0.08)
        pad_x, pad_y = 1.8, (1.2 if small else 1.6)
        ix, iy, iw = b["x"] + pad_x, b["y"] + pad_y, b["w"] - 2 * pad_x
        head_pt = PT_HEAD * (0.7 if small else 1)
        chip = 2.2 if small else 2.9
        tx = ix
        if b["num"]:
            c = sh.add_shape(MSO_SHAPE.OVAL, Cm(ix * s), Cm(iy * s), Cm(chip * s), Cm(chip * s))
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(INK if b.get("accent") else GOLD)
            c.line.fill.background()
            c.shadow.inherit = False
            tf = c.text_frame
            for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                setattr(tf, m, 0)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = b["num"]
            r.font.size = Pt(PT_HEAD * (0.45 if small else 0.62) * s)
            r.font.bold = True
            r.font.name = "Be Vietnam Pro"
            r.font.color.rgb = rgb(RAISED)
            tx = ix + chip + 0.8
        text(tx, iy, b["x"] + b["w"] - pad_x - tx, chip,
             [(b["title"], head_pt, MOSS if b.get("muted") else INK, "Be Vietnam Pro", True, False)],
             anchor=MSO_ANCHOR.MIDDLE)
        y = iy + chip + (0.6 if small else 1.0)
        bottom = b["y"] + b["h"] - pad_y
        hint_pt = PT_SMALL if small else PT_BODY
        if b.get("image"):
            text(ix, y, iw, b["hint_h"], [(b["hint"], hint_pt, SLATE, "IBM Plex Sans", False, True)])
            y += b["hint_h"] + 1.0
            ih = b.get("image_h") or (bottom - y)
            slot = rect(ix, y, iw, ih, PAPER, SLATE, 0.12, dash=True)
            tf = slot.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = f"[ {b['image']} ]"
            r.font.size = Pt(PT_SMALL * s)
            r.font.color.rgb = rgb(SLATE)
            r.font.name = "IBM Plex Sans"
            if b.get("image_h"):
                y += ih + 1.0
                text(ix, y, iw, bottom - y,
                     [("Chú thích ảnh và 3–4 dòng diễn giải.", PT_SMALL, SLATE, "IBM Plex Sans", False, True)])
        else:
            text(ix, y, iw, bottom - y, [(b["hint"], hint_pt, SLATE, "IBM Plex Sans", False, True)])

    notes = slide.notes_slide.notes_text_frame
    notes.text = ("Tỉ lệ 1:2 — slide 100 x 75 cm, in phóng 200% ra 200 x 150 cm.\n"
                  "Nếp gấp ở 25 cm và 75 cm trên slide (50 cm và 150 cm khi in).\n"
                  "Không đặt: " + ", ".join(BANNED) + ".")
    prs.save(path)


# ================================================================ ảnh gập
def folded_html(png_name):
    # Mỗi lá lấy đúng phần của ảnh phẳng; hai lá bên xoay quanh nếp gấp.
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
body {{ margin:0; width:1600px; height:1100px; background:#E6E7E2; font-family:sans-serif; overflow:hidden; }}
.scene {{ position:absolute; left:800px; top:140px; width:0; height:0; perspective:2800px;
         perspective-origin:0 250px; }}
.board {{ position:absolute; left:-500px; top:0; width:1000px; height:1500px; transform-style:preserve-3d;
         transform-origin:50% 0; transform:scale(.5); }}
.p {{ position:absolute; top:0; height:1500px; background:url({png_name}) no-repeat;
     background-size:2000px 1500px; box-shadow:0 0 0 1px rgba(0,0,0,.08); }}
.c {{ left:0; width:1000px; background-position:-500px 0; }}
.l {{ left:-500px; width:500px; background-position:0 0; transform-origin:right center;
     transform:rotateY(62deg); filter:brightness(.93); }}
.r {{ left:1000px; width:500px; background-position:-1500px 0; transform-origin:left center;
     transform:rotateY(-62deg); filter:brightness(.93); }}
.floor {{ position:absolute; left:-300px; top:1480px; width:1600px; height:60px;
         background:radial-gradient(ellipse at center, rgba(0,0,0,.28), transparent 70%); }}
.cap {{ position:fixed; bottom:24px; left:0; right:0; text-align:center; color:#5B6570; font-size:22px; }}
</style></head><body>
<div class="scene"><div class="board"><div class="floor"></div>
<div class="p l"></div><div class="p c"></div><div class="p r"></div></div></div>
<div class="cap">Minh hoạ khi gập — hai lá bên gập vào 90° khi trưng bày thật: rộng 100 cm, sâu 50 cm, cao 150 cm</div>
</body></html>"""


def _chromium():
    """Dùng Chromium có sẵn trên máy nếu bản Playwright chưa tải trình duyệt riêng."""
    import glob
    for pat in ("/opt/pw-browsers/chromium", "/opt/pw-browsers/chromium-*/chrome-linux/chrome"):
        for c in sorted(glob.glob(pat)):
            if Path(c).is_file():
                return c
    return None


def main():
    from playwright.sync_api import sync_playwright

    (HERE / "poster_khung.html").write_text(build_html(), encoding="utf-8")
    (HERE / "_huong_dan.html").write_text(build_html(guide=True), encoding="utf-8")

    with sync_playwright() as p:
        br = p.chromium.launch(executable_path=_chromium())
        pg = br.new_page()
        for src, pdf in [("poster_khung.html", "poster_khung_200x150cm.pdf"),
                         ("_huong_dan.html", "poster_huong_dan.pdf")]:
            pg.goto((HERE / src).as_uri())
            pg.evaluate("document.fonts.ready")
            pg.pdf(path=str(HERE / pdf), width=f"{W}cm", height=f"{H}cm",
                   print_background=True, prefer_css_page_size=True)
        # ảnh xem trước: 10 px/cm (1 cm CSS = 37,8 px → thu theo tỉ lệ)
        pg2 = br.new_page(viewport={"width": 2000, "height": 1500})
        pg2.goto((HERE / "poster_khung.html").as_uri())
        pg2.evaluate("document.fonts.ready")
        z = 2000 / (W * 96 / 2.54)
        pg2.evaluate(f"document.body.style.zoom='{z}'")
        pg2.screenshot(path=str(HERE / "xem_truoc_phang.png"), clip={"x": 0, "y": 0, "width": 2000, "height": 1500})
        (HERE / "_gap.html").write_text(folded_html("xem_truoc_phang.png"), encoding="utf-8")
        pg3 = br.new_page(viewport={"width": 1600, "height": 1100})
        pg3.goto((HERE / "_gap.html").as_uri())
        pg3.wait_for_timeout(300)
        pg3.screenshot(path=str(HERE / "xem_truoc_gap.png"))
        br.close()

    for tmp in ("_huong_dan.html", "_gap.html"):
        (HERE / tmp).unlink()
    build_pptx(HERE / "poster_khung_1-2.pptx")
    print("Đã sinh tệp trong", HERE)


if __name__ == "__main__":
    main()
